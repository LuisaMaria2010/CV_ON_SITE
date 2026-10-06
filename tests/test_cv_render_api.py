"""Tests per POST /api/cv/render e l'anteprima HTML del CV MC (services/cv_mcflash_preview.py)."""
from __future__ import annotations

import asyncio
import base64
import io
import json

import azure.functions as func
from docx import Document

import functions.cv_render_api as api
from services.cv_mcflash_preview import check_context, render_html
from services.cv_mcflash_renderer import build_context

DATA = {
    "first_name": "Mario", "last_name": "Rossi", "full_name": "Mario Rossi", "trittico": "MRO",
    "email": "mario.rossi@example.com", "phone": "+39 333 1234567", "age": 39,
    "birth_date": "1987-05-12", "location": "Torino", "catalog_role": None, "role": "Senior Engineer",
    "experiences": [{
        "start_date": "2020-01-01", "end_date": None, "ongoing": True, "company": "ACME <Corp> & Co",
        "role": "Senior Engineer", "description": "Sviluppo backend\nCode review",
        "technologies": ["Python", "Azure"],
        "projects": [{"name": "Portale clienti", "activities": "API REST", "technologies": ["FastAPI"]}],
    }],
    "education": [{"year": "2012", "title": "MSc Computer Science", "institution": "University X"}],
    "trainings": [{"year": None, "title": "AZ-204", "provider": "Microsoft"}],
    "cv_skills": [{"name": "Python", "level": "Ottima conoscenza"}, {"name": "Docker", "level": None}],
    "languages": [{"lang": "inglese", "level": "Avanzato"}],
}


def _call(body):
    req = func.HttpRequest(method="POST", url="/api/cv/render", body=json.dumps(body).encode(),
                           headers={"content-type": "application/json"}, params={}, route_params={})
    resp = asyncio.run(api.cv_render(req))
    return resp.status_code, json.loads(resp.get_body())


def test_render_both_formats():
    status, env = _call({"data": DATA})
    assert status == 200
    data = env["data"]
    html = data["html"]
    assert "Candidate Profile" in html and "MRO" in html and "Work Experiences" in html
    assert "ACME &lt;Corp&gt; &amp; Co" in html  # testo escapato
    assert "Tecnologie utilizzate: Python, Azure." in html
    docx = base64.b64decode(data["docx_base64"])
    assert len(docx) == data["docx_bytes"]
    text = "\n".join(p.text for t in Document(io.BytesIO(docx)).tables for r in t.rows for c in r.cells
                     for p in c.paragraphs)
    assert "ACME <Corp> & Co" in text


def test_cv_stays_anonymous():
    _, env = _call({"data": DATA})
    html = env["data"]["html"]
    for personal in ("Mario", "Rossi", "mario.rossi@example.com", "1234567"):
        assert personal not in html


def test_invalid_email_does_not_block_preview():
    status, env = _call({"data": {**DATA, "email": "non-una-email"}, "format": "html"})
    assert status == 200 and "html" in env["data"] and "docx_base64" not in env["data"]


def test_checks_list_what_is_missing():
    _, env = _call({"data": DATA, "format": "html"})
    checks = " ".join(env["data"]["checks"])
    assert "Skill senza livello: Docker" in checks
    assert "progetti senza date (Portale clienti)" in checks
    assert "Trainings senza anno: AZ-204" in checks


def test_bad_requests():
    assert _call({})[0] == 400
    assert _call({"data": DATA, "format": "pdf"})[0] == 400
    status, env = _call({"data": {**DATA, "birth_date": "non una data"}})
    assert status == 400 and "birth_date" in env["error"]


def test_empty_sections_are_omitted():
    context = build_context({"trittico": "ABC"})
    html = render_html(context)
    assert "Work Experiences" not in html and "Languages" not in html
    assert "Nessuna esperienza lavorativa." in check_context(context)
