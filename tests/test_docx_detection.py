"""Un DOCX senza MIME (POST /api/extract) va letto come DOCX, non come testo.

Prima veniva letto come testo: lo ZIP grezzo finiva all'LLM, che inventava i
dati, e il risultato restava in cache.
"""
from __future__ import annotations

import asyncio
import io

import pytest
from docx import Document

import db_data.pipeline as pipeline_mod
from extraction.extract import _is_docx, extract_elements


def _docx_bytes(*paragraphs: str) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_detected_without_mime():
    data = _docx_bytes("Mario Rossi")
    assert _is_docx(None, data)
    assert _is_docx("application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"")
    assert not _is_docx(None, b"testo semplice")
    assert not _is_docx(None, b"PK\x03\x04 non uno zip valido")


def test_docx_text_is_extracted_without_mime():
    elements = extract_elements(_docx_bytes("Mario Rossi", "Senior Engineer at ACME Corp"))
    text = " ".join(e.text for e in elements)
    assert "ACME Corp" in text
    assert "PK" not in text[:10]


def test_raw_zip_text_is_recognised():
    assert pipeline_mod.is_raw_zip_text("PK\x03\x04\x14\x00...")
    assert not pipeline_mod.is_raw_zip_text("Mario Rossi\nSenior Engineer")


class _Cache:
    def __init__(self, text, json_payload):
        self.text, self.json_payload, self.saved_text = text, json_payload, None

    async def get(self, key):
        return self.text

    async def save(self, key, text):
        self.saved_text = text

    async def get_json(self, key):
        return self.json_payload

    async def save_json(self, key, payload):
        self.json_payload = payload


def test_stale_zip_cache_is_ignored(monkeypatch):
    """Testo in cache = ZIP grezzo: si estrae di nuovo e non si usa il JSON LLM in cache."""
    cache = _Cache("PK\x03\x04garbage", {"full_name": "Inventato"})
    pipe = pipeline_mod.CVPipeline(cache)
    seen = {}

    class _Chain:
        cache_signature = "sig"
        schema = dict

        async def extract(self, text):
            seen["text"] = text
            raise RuntimeError("stop dopo l'estrazione del testo")

    monkeypatch.setattr(pipe, "_chain_for", lambda profile: _Chain())
    with pytest.raises(RuntimeError, match="stop"):
        asyncio.run(pipe.process(_docx_bytes("Mario Rossi", "ACME Corp")))
    assert "ACME Corp" in seen["text"]
    assert cache.saved_text and not pipeline_mod.is_raw_zip_text(cache.saved_text)


def _docx_with_table(rows, before="Information", after="Languages") -> bytes:
    doc = Document()
    doc.add_paragraph(before)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    doc.add_paragraph(after)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_tables_are_read_in_document_order():
    """I CV nel modello MC (e molti altri) mettono i contenuti in tabelle."""
    data = _docx_with_table([["Città", "Bologna"], ["Ruolo", "Software Developer"]])
    elements = extract_elements(data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    kinds = [e.element_type for e in elements]
    assert kinds == ["paragraph", "table", "paragraph"]
    assert elements[1].rows == [["Città", "Bologna"], ["Ruolo", "Software Developer"]]


def test_merged_cells_count_once():
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "ACME Corp"
    table.cell(0, 2).text = "2020"
    buf = io.BytesIO()
    doc.save(buf)
    rows = extract_elements(buf.getvalue())[0].rows
    assert rows == [["ACME Corp", "2020"]]


def test_two_mc_cvs_have_different_text():
    """Prima tutti i CV MC avevano lo stesso testo (solo i titoli) e quindi lo stesso hash."""
    from services.document_processor import DocumentProcessor

    a = _docx_with_table([["Città", "Bologna"]])
    b = _docx_with_table([["Città", "Torino"]])
    proc = DocumentProcessor()
    ha = proc.process(a, mime_type=None, filename="a.docx", source_path="/a.docx")["content_hash"]
    hb = proc.process(b, mime_type=None, filename="b.docx", source_path="/b.docx")["content_hash"]
    assert ha != hb
