"""
Costruisce il modello CV MCFlash ("Candidate Profile" MC Engineering) con i
segnaposto docxtpl, ricalcando il layout dei CV MC reali.

Scrive:
- templates/cv_mcflash.docx  -> usato da services/cv_mcflash_renderer.py
- templates/cv_mcflash.dotx  -> stesso documento come modello di Word

Immagini (estratte da un CV MC reale): templates/assets/
- mc_logo.png, mc_vitruviano.png (intestazione), mc_footer_sedi.png (footer)

Segnaposto (contesto prodotto da cv_mcflash_renderer.build_context):
    codice, data_nascita, citta, ruolo
    formazione[]  {anno, titolo, istituto}
    training[]    {anno, titolo, ente}
    esperienze[]  {periodo, azienda, ruolo, righe[], tecnologie,
                   progetti[] {periodo, nome, ruolo, righe[], tecnologie}}
    skills[]      {nome, livello}
    lingue[]      {lingua, livello}
Le sezioni con lista vuota spariscono dal CV.

Uso:
    python scripts/build_mcflash_template.py
"""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "templates"
ASSETS = TEMPLATES / "assets"

BLUE = RGBColor(0x1F, 0x38, 0x64)
DARK = RGBColor(0x2C, 0x2C, 0x2C)
BAR_FILL = "CCCCCC"
TITLE_FONT = "Bahnschrift Light"
BODY_FONT = "Calibri"

# Pagina A4 e margini ricavati dal PDF di riferimento (punti -> cm).
PAGE_W, PAGE_H = Cm(21.0), Cm(29.7)
MARGIN_LEFT, MARGIN_RIGHT = Cm(2.2), Cm(1.75)
MARGIN_TOP, MARGIN_BOTTOM = Cm(4.6), Cm(2.9)
CONTENT_W = PAGE_W - MARGIN_LEFT - MARGIN_RIGHT
LABEL_W = Cm(4.0)
VALUE_W = CONTENT_W - LABEL_W


# ---------------------------------------------------------------- helpers XML

def _set_font(run, name: str, size: float, *, bold=False, italic=False, color: RGBColor | None = None,
              small_caps=False, underline=False):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.small_caps = small_caps
    run.font.underline = underline
    if color is not None:
        run.font.color.rgb = color


def _spacing(paragraph, before=0.0, after=0.0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def _shade_paragraph(paragraph, fill: str):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _cant_split(row):
    """La riga non si spezza tra due pagine (skill, lingue, titoli...)."""
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _fixed_widths(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def _para(cell_or_doc, text="", *, font=BODY_FONT, size=11.0, bold=False, italic=False, color=None,
          align=None, before=0.0, after=0.0, first=False):
    """Aggiunge un paragrafo (o usa il primo della cella, se first=True)."""
    if first and hasattr(cell_or_doc, "paragraphs") and cell_or_doc.paragraphs and not cell_or_doc.paragraphs[0].text:
        p = cell_or_doc.paragraphs[0]
    else:
        p = cell_or_doc.add_paragraph()
    if text:
        _set_font(p.add_run(text), font, size, bold=bold, italic=italic, color=color)
    if align is not None:
        p.alignment = align
    _spacing(p, before, after)
    return p


def _tag_paragraph(container, tag: str):
    """Paragrafo che contiene solo un tag docxtpl {%p ... %}: sparisce nel rendering."""
    p = container.add_paragraph()
    _set_font(p.add_run(tag), BODY_FONT, 2)
    _spacing(p)
    return p


# ---------------------------------------------------------------- blocchi

def _section_bar(doc, title: str):
    p = doc.add_paragraph()
    _set_font(p.add_run(title), TITLE_FONT, 14, bold=True, color=BLUE, small_caps=True)
    _shade_paragraph(p, BAR_FILL)
    _spacing(p, before=10, after=6)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    return p


def _two_col_table(doc, rows: int):
    table = doc.add_table(rows=rows, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _no_borders(table)
    _fixed_widths(table, (LABEL_W, VALUE_W))
    for row in table.rows:
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    return table


def _loop_table(doc, loop: str, fill_row, *, keep_rows_whole=True):
    """Tabella a due colonne con una riga ripetuta: {%tr for %} / riga / {%tr endfor %}."""
    table = _two_col_table(doc, 3)
    _para(table.rows[0].cells[0], f"{{%tr for {loop} %}}", size=2, first=True)
    fill_row(table.rows[1].cells[0], table.rows[1].cells[1])
    if keep_rows_whole:
        _cant_split(table.rows[1])
    _para(table.rows[2].cells[0], "{%tr endfor %}", size=2, first=True)
    return table


def _label(cell, text):
    return _para(cell, text, italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, first=True, after=2)


def _header(section):
    header = section.header
    header.is_linked_to_previous = False
    table = header.add_table(rows=1, cols=2, width=CONTENT_W)
    _no_borders(table)
    _fixed_widths(table, (Cm(12.0), CONTENT_W - Cm(12.0)))
    logo_cell, figure_cell = table.rows[0].cells
    logo_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    p = logo_cell.paragraphs[0]
    p.add_run().add_picture(str(ASSETS / "mc_logo.png"), width=Cm(7.2))
    _spacing(p)
    p = figure_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(str(ASSETS / "mc_vitruviano.png"), width=Cm(2.95))
    _spacing(p)
    # header.add_table lascia un paragrafo vuoto prima della tabella: lo togliamo.
    first = header.paragraphs[0]
    if not first.text:
        first._p.getparent().remove(first._p)

    title = header.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(title.add_run("Candidate Profile"), TITLE_FONT, 18, bold=True, color=BLUE, small_caps=True,
              underline=True)
    _spacing(title, before=2, after=0)


def _footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    # Immagine a tutta pagina: rientri negativi pari ai margini.
    p.paragraph_format.left_indent = -MARGIN_LEFT
    p.paragraph_format.right_indent = -MARGIN_RIGHT
    _spacing(p)
    p.add_run().add_picture(str(ASSETS / "mc_footer_sedi.png"), width=PAGE_W)


# ---------------------------------------------------------------- documento

def build_document() -> Document:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(0)

    section = doc.sections[0]
    section.page_width, section.page_height = PAGE_W, PAGE_H
    section.left_margin, section.right_margin = MARGIN_LEFT, MARGIN_RIGHT
    section.top_margin, section.bottom_margin = MARGIN_TOP, MARGIN_BOTTOM
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0)
    _header(section)
    _footer(section)

    # --- Information
    _section_bar(doc, "Information")
    info = _two_col_table(doc, 4)
    for row, (label, value, bold, size) in zip(info.rows, (
        ("Nome", "{{ codice }}", True, 14),
        ("Data di Nascita", "{{ data_nascita }}", False, 11),
        ("Città", "{{ citta }}", False, 11),
        ("Ruolo", "{{ ruolo }}", True, 11),
    )):
        _label(row.cells[0], label)
        _para(row.cells[1], value, bold=bold, size=size, first=True, after=2)

    # --- Education
    _tag_paragraph(doc, "{%p if formazione %}")
    _section_bar(doc, "Education")

    def _education_row(left, right):
        _label(left, "{{ e.anno }}")
        _para(right, "{{ e.titolo }}", bold=True, size=10.5, color=DARK, first=True)
        _para(right, "{{ e.istituto }}", after=6)

    _loop_table(doc, "e in formazione", _education_row)
    _tag_paragraph(doc, "{%p endif %}")

    # --- Training (opzionale)
    _tag_paragraph(doc, "{%p if training %}")
    _section_bar(doc, "Trainings")

    def _training_row(left, right):
        _label(left, "{{ t.anno }}")
        _para(right, "{{ t.titolo }}", bold=True, size=10.5, color=DARK, first=True)
        _tag_paragraph(right, "{%p if t.ente %}")
        _para(right, "{{ t.ente }}")
        _tag_paragraph(right, "{%p endif %}")
        _para(right, "", size=4, after=4)

    _loop_table(doc, "t in training", _training_row)
    _tag_paragraph(doc, "{%p endif %}")

    # --- Work Experiences
    _tag_paragraph(doc, "{%p if esperienze %}")
    _section_bar(doc, "Work Experiences")

    def _activities(cell, owner: str):
        """Righe di attivita' + "Tecnologie utilizzate" di un lavoro o di un progetto."""
        _tag_paragraph(cell, f"{{%p for riga in {owner}.righe %}}")
        _para(cell, "{{ riga }}", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        _tag_paragraph(cell, "{%p endfor %}")
        _tag_paragraph(cell, f"{{%p if {owner}.tecnologie %}}")
        _para(cell, f"Tecnologie utilizzate: {{{{ {owner}.tecnologie }}}}", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        _tag_paragraph(cell, "{%p endif %}")

    def _experience_row(left, right):
        _label(left, "{{ x.periodo }}")
        _para(right, "{{ x.azienda }}", bold=True, first=True, after=2)
        _tag_paragraph(right, "{%p if x.ruolo %}")
        _para(right, "Ruolo: {{ x.ruolo }}", after=1)
        _tag_paragraph(right, "{%p endif %}")
        _activities(right, "x")
        # Progetti dentro lo stesso posto di lavoro, come nei CV MC.
        _tag_paragraph(right, "{%p for pr in x.progetti %}")
        # Spazio tra un progetto e l'altro, non subito sotto il ruolo.
        _tag_paragraph(right, "{%p if x.righe or x.tecnologie or not loop.first %}")
        _para(right, "", size=6, after=4)
        _tag_paragraph(right, "{%p endif %}")
        _tag_paragraph(right, "{%p if pr.periodo %}")
        _para(right, "{{ pr.periodo }}")
        _tag_paragraph(right, "{%p endif %}")
        _para(right, "{{ pr.nome }}")
        _tag_paragraph(right, "{%p if pr.ruolo %}")
        _para(right, "Ruolo: {{ pr.ruolo }}")
        _tag_paragraph(right, "{%p endif %}")
        _activities(right, "pr")
        _tag_paragraph(right, "{%p endfor %}")
        _para(right, "", size=6, after=8)

    # Le esperienze lunghe devono poter continuare alla pagina dopo.
    _loop_table(doc, "x in esperienze", _experience_row, keep_rows_whole=False)
    _tag_paragraph(doc, "{%p endif %}")

    # --- Synthesis of Technical Skills
    _tag_paragraph(doc, "{%p if skills %}")
    _section_bar(doc, "Synthesis of Technical Skills")

    def _skill_row(left, right):
        _label(left, "{{ s.nome }}")
        _para(right, "{{ s.livello }}", first=True, after=2)

    _loop_table(doc, "s in skills", _skill_row)
    _tag_paragraph(doc, "{%p endif %}")

    # --- Languages
    _tag_paragraph(doc, "{%p if lingue %}")
    _section_bar(doc, "Languages")

    def _language_row(left, right):
        _label(left, "{{ l.lingua }}")
        _para(right, "{{ l.livello }}", bold=True, first=True, after=2)

    _loop_table(doc, "l in lingue", _language_row)
    _tag_paragraph(doc, "{%p endif %}")

    doc.core_properties.title = "Candidate Profile"
    doc.core_properties.author = "MC Engineering"
    return doc


def _as_dotx(docx_bytes: bytes) -> bytes:
    """Stesso pacchetto con il content type del modello di Word."""
    src = zipfile.ZipFile(io.BytesIO(docx_bytes))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"wordprocessingml.document.main+xml",
                    b"wordprocessingml.template.main+xml",
                )
            dst.writestr(item, data)
    return out.getvalue()


def main() -> None:
    TEMPLATES.mkdir(exist_ok=True)
    buf = io.BytesIO()
    build_document().save(buf)
    docx_bytes = buf.getvalue()
    (TEMPLATES / "cv_mcflash.docx").write_bytes(docx_bytes)
    (TEMPLATES / "cv_mcflash.dotx").write_bytes(_as_dotx(docx_bytes))
    print(f"Scritti {TEMPLATES / 'cv_mcflash.docx'} e {TEMPLATES / 'cv_mcflash.dotx'}")


if __name__ == "__main__":
    main()
