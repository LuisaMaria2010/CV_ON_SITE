"""CV nel modello MCFlash: contesto dai dati estratti e compilazione del .docx."""
from __future__ import annotations

import io
from datetime import date

import pytest
from docx import Document

from core.schema import CVExtraction
from services.cv_mcflash_renderer import TEMPLATE_PATH, build_context, render_docx

pytestmark = pytest.mark.skipif(not TEMPLATE_PATH.exists(), reason="eseguire scripts/build_mcflash_template.py")


def _cv(**overrides) -> CVExtraction:
    data = dict(
        full_name="Davide Carbone",
        first_name="Davide",
        last_name="Carbone",
        email="davide@example.com",
        phone="+39 333 0000000",
        trittico="DCA",
        birth_date=date(1990, 9, 1),
        location="Roma",
        role="Business Intelligence Developer",
        experiences=[
            {"start_date": date(2023, 6, 1), "ongoing": True, "company": "Riservato", "role": "BI Developer",
             "description": "Report Power BI per R&D <interno>\nMigrazione Qlik"},
            {"start_date": date(2021, 12, 1), "end_date": date(2023, 6, 1), "company": "Lutech", "role": "Data Analyst"},
        ],
        education=[{"year": "2021", "title": "Laurea Magistrale in Statistica", "institution": "Unical"}],
        languages=[{"lang": "inglese", "level": "Intermedio"}],
        skills=[{"name": "power bi", "level": "Conoscenza base"}],
        cv_skills=[{"name": "Power BI", "level": "Conoscenza base"}],
    )
    data.update(overrides)
    return CVExtraction(**data)


def _text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts += [p.text for p in cell.paragraphs]
    return "\n".join(parts)


def test_context_formats_mc_conventions():
    ctx = build_context(_cv())
    assert ctx["codice"] == "DCA"
    assert ctx["data_nascita"] == "01/09/1990"
    assert [x["periodo"] for x in ctx["esperienze"]] == ["Da 06/2023 ad Oggi", "Da 12/2021 a 06/2023"]
    assert ctx["esperienze"][0]["righe"] == ["Attività: Report Power BI per R&D <interno>", "Migrazione Qlik"]
    assert ctx["skills"] == [{"nome": "Power BI", "livello": "Conoscenza base"}]
    assert ctx["lingue"] == [{"lingua": "Inglese", "livello": "Intermedio"}]


def test_context_accepts_edited_dict_from_dashboard():
    edited = _cv().model_dump(mode="json")
    edited["role"] = "BI Lead"
    assert build_context(edited)["ruolo"] == "BI Lead"


def test_context_falls_back_to_base_profile_fields():
    ctx = build_context(_cv(education=[], cv_skills=[], education_titles=["Laurea in Fisica"]))
    assert ctx["formazione"] == [{"anno": "", "titolo": "Laurea in Fisica", "istituto": ""}]
    assert ctx["skills"] == [{"nome": "power bi", "livello": "Conoscenza base"}]


def test_render_fills_template_and_keeps_cv_anonymous():
    text = _text(render_docx(build_context(_cv())))
    for expected in ("DCA", "01/09/1990", "Riservato", "Da 06/2023 ad Oggi", "R&D <interno>",
                     "Laurea Magistrale in Statistica", "Power BI", "Inglese"):
        assert expected in text
    assert "{{" not in text and "{%" not in text
    # Il CV MC e' anonimo.
    for secret in ("Davide", "Carbone", "davide@example.com", "333 0000000"):
        assert secret not in text


def _cv_with_projects() -> CVExtraction:
    return _cv(
        experiences=[{
            "start_date": date(2023, 9, 1), "ongoing": True, "company": "MC Engineering", "role": "Data Scientist",
            "projects": [
                {"start_date": date(2024, 9, 1), "ongoing": True, "name": "Progetto Angelini",
                 "role": "Data Scientist", "activities": "Chatbot ASKAI su GPT-4",
                 "technologies": ["Azure OpenAI", "Python"]},
                {"start_date": date(2024, 2, 1), "end_date": date(2024, 9, 1), "name": "Progetto MedioBanca",
                 "role": "Tech Lead", "activities": "Attività: interfaccia web per riconoscimento documentale"},
            ],
        }],
        trainings=[{"year": "2024", "title": "Advanced Data Analytics using Python", "provider": "Udemy"},
                   {"year": "2023", "title": "SQL class"}],
    )


def test_context_projects_and_trainings():
    ctx = build_context(_cv_with_projects())
    job = ctx["esperienze"][0]
    assert job["righe"] == [] and job["tecnologie"] == ""
    first, second = job["progetti"]
    assert first == {
        "periodo": "Da 09/2024 ad Oggi", "nome": "Progetto Angelini", "ruolo": "",
        "righe": ["Attività: Chatbot ASKAI su GPT-4"], "tecnologie": "Azure OpenAI, Python.",
    }
    # Ruolo diverso da quello del lavoro: si mostra; "Attività:" non si duplica.
    assert second["ruolo"] == "Tech Lead"
    assert second["righe"] == ["Attività: interfaccia web per riconoscimento documentale"]
    assert second["periodo"] == "Da 02/2024 a 09/2024"
    assert ctx["training"] == [
        {"anno": "2024", "titolo": "Advanced Data Analytics using Python", "ente": "Udemy"},
        {"anno": "2023", "titolo": "SQL class", "ente": ""},
    ]


def test_render_projects_and_trainings():
    text = _text(render_docx(build_context(_cv_with_projects())))
    for expected in ("MC Engineering", "Ruolo: Data Scientist", "Da 09/2024 ad Oggi", "Progetto Angelini",
                     "Attività: Chatbot ASKAI su GPT-4", "Tecnologie utilizzate: Azure OpenAI, Python.",
                     "Progetto MedioBanca", "Ruolo: Tech Lead", "Trainings", "Udemy"):
        assert expected in text
    assert "{{" not in text and "{%" not in text
    assert "None" not in text


def test_render_omits_empty_sections():
    text = _text(render_docx(build_context(_cv(certifications=[], languages=[]))))
    assert "Training" not in text
    assert "Languages" not in text
    assert "Work Experiences" in text
